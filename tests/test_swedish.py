"""End-to-end tests for searching in Swedish and for what a phone shows first.

    python3 build.py && python3 tests/test_swedish.py

Like test_app.py, JobTech's services are replaced with fixed answers. One county search for "undersköterska" in Skåne,
with "natt" left out, covers what the audit of 25 October 2026 found: a word left out also leaves out compounds
("Nattpatrull"), keywords count plural forms and compounds ("Undersköterskor", "Specialistundersköterska"), the same ad
from another job site is dropped even when the employer is written another way, the snippet shows where the keyword is,
"Ads with every keyword" switched on by the app doesn't stick, filters that hide roles say so, the Excel file, and on a
phone the folded search form, the filters behind a button and the results on the first screen."""
import asyncio, io, json, pathlib, sys, zipfile
from urllib.parse import urlparse, parse_qs
from playwright.async_api import async_playwright

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from test_app import ok, ad, MAL, check, passed, EVENT_HOSTS, no_events, history

ROOT = pathlib.Path(__file__).resolve().parent.parent
PAGE = (ROOT / 'index.html').as_uri()
HELS = {**MAL, 'municipality': 'Helsingborg', 'municipality_concept_id': 'mHBG', 'city': 'Helsingborg', 'coordinates': [12.69, 56.05]}
LUND = {**MAL, 'municipality': 'Lund', 'municipality_concept_id': 'mLUN', 'city': 'Lund', 'coordinates': [13.19, 55.70]}
INTRO = 'Hos oss får du kollegor som bryr sig och en arbetsplats med fri parkering, laddstolpar och goda förbindelser med kollektivtrafik. ' * 3
CARE = [ad('401', 'Undersköterska till hemtjänst', 'Lunds kommun', '2120001090', LUND, INTRO + 'Vi söker en undersköterska till hemtjänsten i Lund.', field='Hälso- och sjukvård'),
        ad('402', 'Specialistundersköterska demens', 'Svalövs kommun', '2120001009', MAL, 'Du är specialistundersköterska med demensinriktning.', field='Hälso- och sjukvård'),
        ad('403', 'Undersköterskor till Hemgården', 'Hässleholms kommun', '2120000951', MAL, 'Vi söker undersköterskor dag och kväll.', field='Hälso- och sjukvård'),
        ad('404', 'Undersköterska till Hemtjänstens Nattpatrull', 'Kävlinge kommun', '2120000035', MAL, 'Undersköterska till nattpatrullen.', field='Hälso- och sjukvård'),
        ad('405', 'Undersköterska till Lundhagens vårdboende', 'Helsingborgs kommun', '2120001165', HELS, 'Undersköterska till vårdboendet.', field='Hälso- och sjukvård'),
        ad('406', 'Undersköterska, Clean Care', 'Clean Care AB', '5560000406', MAL, 'Undersköterska och clean room.', field='Hälso- och sjukvård')]
# Other job sites: the same Lundhagen ad with the employer written another way, and one that is only there
LINKS = [{"id": "s1", "headline": "Undersköterska till Lundhagens vårdboende", "brief": "Undersköterska till vårdboendet.", "occupation_field": {"concept_id": "h", "label": "Hälso- och sjukvård"},
          "employer": {"name": "Helsingborgs stad"}, "workplace_addresses": [{"municipality": "Helsingborg", "municipality_concept_id": "mHBG", "region_concept_id": "rSK", "country": "Sverige"}],
          "publication_date": "2026-10-20T00:00:00", "source_links": [{"label": "ledigajobb.se", "url": "https://ledigajobb.se/s1"}]},
         {"id": "s2", "headline": "Undersköterska till Vardaga", "brief": "Undersköterska sökes.", "occupation_field": {"concept_id": "h", "label": "Hälso- och sjukvård"},
          "employer": {"name": "Vardaga"}, "workplace_addresses": [{"municipality": "Malmö", "municipality_concept_id": "mMAL", "region_concept_id": "rSK", "country": "Sverige"}],
          "publication_date": "2026-10-20T00:00:00", "source_links": [{"label": "thehub.io", "url": "https://thehub.io/s2"}]}]
LEAN = [ad('501', 'Cleanroom operator', 'Medtech Skåne AB', '5560000501', MAL, 'Arbete i cleanroom med kvalitetssystem.', field='Tillverkning'),
        ad('502', 'Lean coach', 'Factory Skåne AB', '5560000502', MAL, 'Du coachar team i lean.', field='Tillverkning')]

async def jobsearch(route):
    q = parse_qs(urlparse(route.request.url).query)
    term = q.get('q', [''])[0].strip('"').lower()
    if q.get('qfields') == ['location']:
        hits = [MAL] * 3 if term == 'skåne' else []
        return await route.fulfill(**ok({"total": {"value": len(hits)}, "hits": [{"id": f"g{i}", "workplace_address": w} for i, w in enumerate(hits)]}))
    if q.get('limit') == ['0']: return await route.fulfill(**ok({"total": {"value": 0 if ' ' in term else 50}, "hits": []}))   # no ad uses the phrase
    hits = CARE if term == 'undersköterska' else LEAN if term in ('lean', 'cleanroom') else []
    await route.fulfill(**ok({"total": {"value": len(hits)}, "hits": hits}))
async def joblinks(route):
    term = parse_qs(urlparse(route.request.url).query).get('q', [''])[0].strip('"').lower()
    hits = LINKS if term == 'undersköterska' else []
    await route.fulfill(**ok({"total": {"value": len(hits)}, "hits": hits}))

async def clear(pg, box):   # each × redraws the tags, so they are removed one at a time
    await pg.evaluate("box => { let b; while ((b = document.querySelector('#' + box + ' .tok button'))) b.click(); }", box)
async def search(pg):
    await pg.evaluate("document.getElementById('msg').textContent=''")
    await pg.click('#go'); await pg.wait_for_function("document.getElementById('msg').textContent.length > 0", timeout=20000)

async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        ctx = await browser.new_context(locale='en-US', viewport={'width': 1280, 'height': 1000}, accept_downloads=True)
        await ctx.route('https://jobsearch.api.jobtechdev.se/**', jobsearch)
        await ctx.route('https://links.api.jobtechdev.se/**', joblinks)
        await ctx.route('https://historical.api.jobtechdev.se/**', history)
        await ctx.route(EVENT_HOSTS, no_events)
        pg = await ctx.new_page(); errors = []; pg.on('pageerror', lambda e: errors.append(str(e)))
        await pg.goto(PAGE)

        await pg.fill('#place', 'Skåne län')
        await pg.fill('#kw-input', 'undersköterska'); await pg.press('#kw-input', 'Enter')
        await pg.fill('#ex-input', 'natt'); await pg.press('#ex-input', 'Enter')
        await search(pg)
        titles = await pg.evaluate("[...document.querySelectorAll('article.role h3, article.role .role-title, article.role a.title')].map(e => e.textContent.trim())")
        ads = await pg.evaluate("JSON.parse(localStorage.getItem('jf-results')).ads")
        by = {a['title']: a for a in ads}
        check('a word left out also leaves out the words it starts ("natt" and "Nattpatrull")', not any('Nattpatrull' in a['title'] for a in ads), [a['title'] for a in ads])
        check('plural forms count as the keyword', by['Undersköterskor till Hemgården']['keys'] == ['undersköterska'], by['Undersköterskor till Hemgården'])
        check('a compound that ends in the keyword counts', by['Specialistundersköterska demens']['keys'] == ['undersköterska'], by['Specialistundersköterska demens'])
        check('the same ad on another job site is dropped when the employer is written another way',
              sum(a['title'] == 'Undersköterska till Lundhagens vårdboende' for a in ads) == 1 and any(a['employer'] == 'Vardaga' for a in ads), [(a['title'], a['employer']) for a in ads])
        lund = by['Undersköterska till hemtjänst']['snippet']
        check('the snippet starts near the keyword, not with the employer\'s intro', lund.startswith('…') and 'undersköterska till hemtjänsten' in lund and lund.count('fri parkering') <= 1, lund[:200])

        # Every keyword, switched on by splitting a phrase, holds for that search only
        await pg.fill('#kw-input', 'undersköterska hemtjänst nattpatrull'); await pg.press('#kw-input', 'Enter')
        await pg.evaluate("document.querySelectorAll('#kw-box .tok button')[0].click()")   # only the phrase is left
        await pg.wait_for_selector('#kw-hint [data-kwh=split]', timeout=10000)
        await pg.click('#kw-hint [data-kwh=split]')
        check('splitting a phrase shows only ads with every word', await pg.is_checked('#f-all'))
        check('but that is not saved as the person\'s choice', not await pg.evaluate("JSON.parse(localStorage.getItem('jf-prefs') || '{}').all"))
        await clear(pg, 'kw-box')
        await pg.fill('#kw-input', 'undersköterska'); await pg.press('#kw-input', 'Enter')
        await search(pg)
        check('a search for other keywords puts the choice back', not await pg.is_checked('#f-all'))

        # A filter that hides roles says which one, with a way to clear it
        await pg.check('#f-direct'); await pg.fill('#q', 'Lund')
        count = await pg.locator('#count').inner_text()
        check('filters that hide roles are named', 'hidden by Filter and Employers hiring directly' in count and 'Clear filters' in count, count)
        await pg.click('#count [data-act=clear-filters]')
        check('Clear filters shows everything again', await pg.input_value('#q') == '' and not await pg.is_checked('#f-direct') and 'hidden' not in await pg.locator('#count').inner_text())

        # The Excel file: the same rows as the CSV, Swedish letters intact, headings bold and frozen
        async with pg.expect_download() as dl: await pg.click('#export-xlsx')
        d = await dl.value
        data = pathlib.Path(await d.path()).read_bytes()
        z = zipfile.ZipFile(io.BytesIO(data)); sheet = z.read('xl/worksheets/sheet1.xml').decode()
        check('an Excel file with every role and Swedish letters', d.suggested_filename.endswith('.xlsx') and 'Hässleholms kommun' in sheet and sheet.count('<row ') == 1 + len(ads)
              and 'state="frozen"' in sheet and set(z.namelist()) >= {'[Content_Types].xml', 'xl/workbook.xml', 'xl/styles.xml'}, d.suggested_filename)
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data)); ws = wb.active
            check('it opens as a spreadsheet', ws['A1'].value == 'Employer' and ws['A1'].font.b and ws.max_row == 1 + len(ads), [ws['A1'].value, ws.max_row])
        except ImportError: pass

        # A short keyword doesn't match inside a longer word
        await clear(pg, 'kw-box')
        await pg.fill('#kw-input', 'lean'); await pg.press('#kw-input', 'Enter')
        await clear(pg, 'ex-box')
        await search(pg)
        lean = {a['title']: a for a in await pg.evaluate("JSON.parse(localStorage.getItem('jf-results')).ads")}
        check('"lean" matches lean, not cleanroom', lean['Lean coach']['keys'] == ['lean'] and lean['Cleanroom operator']['keys'] == [], {k: v['keys'] for k, v in lean.items()})

        # On a phone: the search folds into a line, the filters sit behind a button, and the first employer is on screen
        await pg.set_viewport_size({'width': 375, 'height': 812})
        await clear(pg, 'kw-box')
        await pg.fill('#kw-input', 'undersköterska'); await pg.press('#kw-input', 'Enter')
        await search(pg)
        check('the form folds into one line after a search', await pg.locator('#search').is_hidden() and 'undersköterska · Skåne län' in await pg.locator('#search-sum').inner_text())
        check('the filters sit behind a button', await pg.locator('#f-all').is_hidden() and await pg.locator('#filters-btn').is_visible())
        top = await pg.evaluate("document.querySelector('#list article.co').getBoundingClientRect().top")
        check('the first employer is on the first screen', top < 812, top)
        await pg.click('#filters-btn')
        check('the button opens the filters', await pg.locator('#f-all').is_visible() and await pg.get_attribute('#filters-btn', 'aria-expanded') == 'true')
        await pg.click('#search-edit')
        check('Edit search opens the form again', await pg.locator('#search').is_visible() and await pg.evaluate("document.activeElement.id") == 'kw-input')
        check('no sideways scroll on a phone', await pg.evaluate('document.documentElement.scrollWidth') <= 375)
        await pg.reload()
        check('coming back on a phone shows the folded form and the results', await pg.locator('#search').is_hidden() and await pg.locator('#list article.co').count() > 0)
        await pg.set_viewport_size({'width': 1280, 'height': 1000})
        check('on a wide screen the form is always there', await pg.locator('#search').is_visible() and await pg.locator('#f-all').is_visible() and await pg.locator('#filters-btn').is_hidden())
        check('no script errors', not errors, errors)
        await browser.close()
    print(f'{len(passed)} checks passed')

if __name__ == '__main__':
    try: asyncio.run(main())
    except AssertionError as e: print('FAILED', e); sys.exit(1)
