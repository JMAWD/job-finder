"""End-to-end tests for step 2 (tailoring an application), with the real server code and a fake AI service.

    python3 build.py && python3 tests/test_tailor.py

Starts server/dev.mjs on port 8787 and tests/fake_ai.mjs on port 8799, builds a copy of the page that points at the local
server, and runs through the panel in a browser: fetching the full ad, the CV from a Word file, access code errors,
generating, the whole tailored CV as a PDF and a Word file, editing a text and checking it again, the Word download, backup,
Swedish and a phone-sized screen."""
import asyncio, hashlib, json, os, pathlib, re, shutil, subprocess, sys, tempfile, time, urllib.request, zipfile
from playwright.async_api import async_playwright

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from test_app import jobsearch, joblinks, history, search, ok, SW, check, passed, EVENT_HOSTS, no_events

ROOT = pathlib.Path(__file__).resolve().parent.parent
SHOTS = ROOT / 'docs' / 'screenshots'
CODE = 'TESTA-CODE2'
SERVER, FAKE = 'http://127.0.0.1:8787', 'http://127.0.0.1:8799'
AD_TEXT = ('Om tjänsten\nVi söker en software engineer som kan FMEA och kvalitet. Du kommer att arbeta med riskanalys, '
           'leverantörer och förbättringar i produktionen tillsammans med vårt team i Göteborg.\n\nKvalifikationer\n- Erfarenhet av FMEA\n- Python')
CV = ('Kvalitetsingenjör med åtta års erfarenhet av FMEA, 8D och leverantörsrevisioner inom fordonsindustrin. '
      'Har lett förbättringsarbete i produktionen och utbildat kollegor i riskanalys. ') * 2
# A whole CV for the tailored CV task; tests/fake_ai.mjs answers with a tailored version that only uses what's here
FULL_CV = '''Eva Lindqvist
Göteborg · eva.lindqvist@example.com · 070-123 45 67

Profil
Kvalitetsingenjör med åtta års erfarenhet av FMEA, 8D och leverantörsrevisioner inom fordonsindustrin. Har lett förbättringsarbete i produktionen och utbildat kollegor i riskanalys.

Erfarenhet
Kvalitetsingenjör, Nordvik Automotive AB, Göteborg
2018 – nu
- Ledde FMEA för bromssystem tillsammans med konstruktion och produktion
- Genomförde leverantörsrevisioner och 8D-utredningar
- Utbildade kollegor i riskanalys

Produktionstekniker, Svea Bearings AB, Göteborg
2016 – 2018
- Ansvarade för förbättringsarbete på monteringslinan

Utbildning
Högskoleingenjör i maskinteknik, Chalmers, 2013 – 2016

Språk
Svenska (modersmål), engelska (flytande)'''

def make_docx(paragraphs, path):
    """A Word file with compressed parts, the way Word itself saves them."""
    num = '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
    body = ''.join(f'<w:p>{num if bullet else ""}<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>' for text, bullet in paragraphs)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr('_rels/.rels', '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr('word/document.xml', f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{body}</w:body></w:document>')

def post(url, data):
    req = urllib.request.Request(url, data=json.dumps(data).encode(), headers={'Content-Type': 'application/json'}, method='POST')
    return json.loads(urllib.request.urlopen(req).read())
last_ai = lambda: json.loads(urllib.request.urlopen(FAKE + '/last').read())

AD_DOWN = {'on': False}
async def ad_or_search(route):
    url = route.request.url
    if '/ad/' in url and AD_DOWN['on']: return await route.abort()
    if '/ad/' in url:
        a = next((x for x in SW if url.rstrip('/').endswith('/' + x['id'])), None)
        return await route.fulfill(**ok({**a, 'description': {'text': AD_TEXT}})) if a else await route.fulfill(status=404, body='{}')
    await jobsearch(route)

async def main():
    tmp = pathlib.Path(tempfile.mkdtemp())
    html = (ROOT / 'index.html').read_text(encoding='utf-8')
    page, off = tmp / 'Job Finder.html', tmp / 'no-server.html'
    page.write_text(re.sub(r'"tailor": "[^"]*"', f'"tailor": "{SERVER}"', html, count=1), encoding='utf-8')
    off.write_text(re.sub(r'"tailor": "[^"]*"', '"tailor": ""', html, count=1), encoding='utf-8')
    env = {**os.environ, 'AI_KEY': 'fake-key', 'ACCESS_CODES': hashlib.sha256(CODE.encode()).hexdigest(), 'ALLOWED_ORIGINS': 'null',
           'AI_BASE_URL': FAKE + '/v1', 'LIMIT_PER_HOUR': '50', 'LIMIT_PER_DAY': '100'}
    procs = [subprocess.Popen(['node', str(ROOT / 'tests' / 'fake_ai.mjs'), '8799']),
             subprocess.Popen(['node', 'dev.mjs'], cwd=ROOT / 'server', env=env, stdout=subprocess.DEVNULL)]
    try:
        for _ in range(50):
            try: urllib.request.urlopen(FAKE + '/last'); break
            except Exception: time.sleep(0.1)
        async with async_playwright() as p:
            browser = await p.chromium.launch()
            ctx = await browser.new_context(locale='en-US', viewport={'width': 1280, 'height': 1000}, accept_downloads=True)
            await ctx.grant_permissions(['clipboard-read', 'clipboard-write'])
            await ctx.route('https://jobsearch.api.jobtechdev.se/**', ad_or_search)
            await ctx.route(EVENT_HOSTS, no_events)
            await ctx.route('https://links.api.jobtechdev.se/**', joblinks)
            await ctx.route('https://historical.api.jobtechdev.se/**', history)
            pg = await ctx.new_page(); errors = []; pg.on('pageerror', lambda e: errors.append(str(e)))

            # Without a server address there is no step 2
            await pg.goto(off.as_uri())
            await pg.fill('#place', 'Göteborg'); await pg.fill('#kw-input', 'software engineer'); await pg.press('#kw-input', 'Enter'); await search(pg)
            check('no tailor buttons without a server', await pg.locator('[data-tailor]').count() == 0)
            check('no Mentor without a server', await pg.locator('#tab-mentor').is_hidden() and await pg.locator('[data-act=goto-mentor]').count() == 0)
            check('no translation buttons without a server', await pg.locator('[data-translate]').count() == 0)

            translation_payloads = []
            async def capture_translation(route):
                payload = route.request.post_data_json
                if isinstance(payload, dict) and payload.get('task') == 'translate': translation_payloads.append(payload)
                await route.continue_()
            await pg.route(SERVER + '/**', capture_translation)
            await pg.goto(page.as_uri()); await search(pg)
            acme = pg.locator('article.co', has_text='Acme Robotics'); await acme.locator('button.exp').click()
            trans_role = acme.locator('article.role[data-ad="101"]')
            translate_btn = trans_role.locator('[data-translate]')
            check('translation button names English as the target', await translate_btn.inner_text() == 'Translate ad to English')
            check('no translation button on an ad already in English', await acme.locator('article.role[data-ad="102"] [data-translate]').count() == 0)
            await translate_btn.click()
            code_input = trans_role.locator('[data-translate-code]')
            check('translation asks for a code but not a CV', await code_input.count() == 1 and 'CV is not needed or sent' in await trans_role.locator('.translation-access').inner_text() and await trans_role.locator('.translation-access textarea').count() == 0)
            await code_input.fill(CODE); await trans_role.locator('[data-translate-go]').click()
            await pg.wait_for_selector('article.role[data-ad="101"] .translation-text', timeout=10000)
            first_translation = await trans_role.locator('.translation-text').inner_text()
            check('full ad translation is shown inline', 'Quality Engineer' in first_translation and 'quality improvements' in first_translation, first_translation)
            check('translation sends only the ad text to the server, without CV or job metadata', len(translation_payloads) == 1 and translation_payloads[0].get('job') == {'text': AD_TEXT} and 'cv' not in translation_payloads[0], repr(translation_payloads))
            first_request = last_ai()['body']['messages'][1]['content']
            check('AI receives the full ad but no CV or separate job metadata', 'Erfarenhet av FMEA' in first_request and 'leverantörsrevisioner' not in first_request and '<cv>' not in first_request and 'Acme Robotics' not in first_request and 'Software Engineer till Acme' not in first_request)
            await pg.evaluate("localStorage.removeItem('jf-code')")
            btn = trans_role.locator('[data-tailor]')
            check('tailor button on each role', await btn.inner_text() == 'Tailor my application')
            await btn.click()
            dlg = pg.locator('#tailor')
            await pg.wait_for_function("document.getElementById('tl-ad') && document.getElementById('tl-ad').value.includes('Kvalifikationer')", timeout=5000)
            check('panel opens with the job title', 'Software Engineer till Acme' in await dlg.locator('h2').inner_text())
            check('full ad fetched from Platsbanken', 'Erfarenhet av FMEA' in await pg.input_value('#tl-ad'))
            check('output language follows the ad', await pg.input_value('#tl-lang') == 'sv')
            check('asks for a code, naming the author', 'Ask John Mawdsley for an access code' in await dlg.inner_text())
            check('before anything is written: a draft to check, not a finished text', "You'll get a draft to read and change, not a finished text." in await dlg.locator('.tl-result').inner_text())

            await dlg.locator('[data-tl=go]').click()
            check('no code: clear message', 'Enter your access code first' in await dlg.locator('.tl-err').inner_text())
            await pg.fill('#tl-code', 'WRONG-CODE9'); await dlg.locator('[data-tl=go]').click()
            check('short CV: clear message', 'at least 200 characters' in await dlg.locator('.tl-err').inner_text())

            # CV from a Word file, with a bullet list
            cvdoc = tmp / 'my-cv.docx'
            make_docx([('Kvalitetsingenjör', False)] + [(line, False) for line in CV.split('. ') if line] + [('FMEA', True), ('8D', True)], cvdoc)
            await pg.set_input_files('#tl-file', str(cvdoc)); await pg.wait_for_timeout(300)
            cv_loaded = await pg.input_value('#tl-cv')
            check('CV read from a compressed Word file', cv_loaded.startswith('Kvalitetsingenjör\n') and '- FMEA\n- 8D' in cv_loaded, cv_loaded[:200])
            check('Word load confirmed', 'Loaded my-cv.docx' in await dlg.locator('#tl-cvmsg').inner_text())

            await dlg.locator('[data-tl=go]').click()
            await pg.wait_for_selector('#tailor .tl-err', timeout=10000)
            check('wrong code refused by the server', "wasn't accepted" in await dlg.locator('.tl-err').inner_text())

            await pg.fill('#tl-code', CODE.lower()); await pg.fill('#tl-extra', 'Håll det kort')
            await dlg.locator('[data-tl=go]').click()
            check('busy state while writing', await dlg.locator('[data-tl=go]').is_disabled())
            await pg.wait_for_selector('#tailor .tl-out', timeout=10000)
            out = dlg.locator('.tl-out')
            check('divider lines left out', '---' not in await out.inner_text())
            check('answer shown with headings and lists', await out.locator('h4').count() == 3 and await out.locator('li').count() == 3 and await out.locator('b').first.inner_text() == 'FMEA')
            note = dlg.locator('.ai-note')
            check('the answer comes with a note that it is an AI draft, above it', await note.count() == 1 and await pg.evaluate("document.querySelector('.ai-note').compareDocumentPosition(document.querySelector('.tl-out')) & 4")
                  and 'Written by AI' == await note.locator('.ai-tag').text_content() and 'This is an estimate from the words in your CV and the ad' in await note.inner_text(), await note.all_inner_texts())
            check('focus lands on the note first', await pg.evaluate("document.activeElement.classList.contains('ai-note')"))
            fc = dlg.locator('.fc')
            check('the answer passed the fact check and says so, between the note and the text', await fc.count() == 1 and 'fc-ok' in (await fc.get_attribute('class'))
                  and 'every employer, job title, school, degree and date in this text is in your CV' in await fc.inner_text()
                  and await pg.evaluate("document.querySelector('.fc').compareDocumentPosition(document.querySelector('.tl-out')) & 4"), await fc.all_inner_texts())
            report = await note.locator('a[data-fb=ai]').get_attribute('href')
            check('the note links to a report about the answer', '%5BAI%5D' in report and 'Part%3A%20Match%20check' in report, report)
            sent = last_ai()['body']
            user = sent['messages'][1]['content']
            check('server sent the ad, the CV and the wish', 'Erfarenhet av FMEA' in user and 'leverantörsrevisioner' in user and 'Håll det kort' in user and '<annons>' in user)
            check('server added its own rules and key', 'Hitta aldrig på' in sent['messages'][0]['content'] and last_ai()['auth'] == 'Bearer fake-key')
            check('job marked Saved after the first result', await acme.locator('article.role[data-ad="101"] select[data-role=status]').input_value() == 'Saved')

            await dlg.locator('[data-tl=copy]').click()
            clip = await pg.evaluate('navigator.clipboard.readText()')
            check('copy gives plain text', clip.replace('\r\n', '\n').startswith('Det som matchar\n• FMEA: åtta') and '**' not in clip and '##' not in clip and '---' not in clip, clip[:80])

            # Cover letter in English, then the Word file
            await dlg.locator('[data-task=letter]').click()
            check('task switch keeps the CV', (await pg.input_value('#tl-cv')) == cv_loaded)
            await pg.select_option('#tl-lang', 'en')
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .tl-out', timeout=10000)
            check('English prompt used', '<job_ad>' in last_ai()['body']['messages'][1]['content'])
            check('a cover letter gets the interview test', "If you couldn't explain something in an interview, take it out." in await dlg.locator('.ai-note').inner_text())
            check('finished tasks get a dot', await dlg.locator('[data-task=match] .tl-dot').count() == 1)

            # The tailored CV is the whole CV, ready to send: shown in full, and downloaded as a PDF or a Word file
            await pg.fill('#tl-cv', FULL_CV)
            await dlg.locator('[data-task=cv]').click(); await pg.select_option('#tl-lang', 'sv')
            check('the CV task promises a whole CV, not tips', 'Your whole CV rewritten for this job' in await dlg.locator('.tl-desc').inner_text()
                  and 'Copy the parts you like' not in await dlg.inner_text())
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .tl-out', timeout=10000)
            ask = last_ai()['body']['messages'][1]['content']
            check('the prompt asks for the whole CV and nothing else', 'skriv hela det nya CV:t, färdigt att skicka' in ask and 'Skriv bara själva CV:t' in ask
                  and 'kan tas bort' not in ask and 'Eva Lindqvist' in ask and last_ai()['body']['max_tokens'] == 4000, ask[:300])
            out = dlg.locator('.tl-out')
            check('the whole CV is shown: the name, both jobs and the education', await out.locator('h3').inner_text() == 'Eva Lindqvist'
                  and await out.locator('h5').count() == 3 and 'Svea Bearings' in await out.inner_text() and 'Chalmers' in await out.inner_text())
            check('the whole CV passes the fact check', await dlg.locator('.fc-ok').count() == 1)
            row = dlg.locator('.tl-result .row').last
            check('PDF comes first, then Word, Edit text and Copy', [await b.get_attribute('data-tl') for b in await row.locator('button').all()] == ['pdf', 'word', 'edit', 'copy'])
            check('no reminder about brackets when there are none', await dlg.locator('.tl-fill').count() == 0)
            async with pg.expect_download() as dl: await dlg.locator('[data-tl=pdf]').click()
            d = await dl.value; cv_pdf = tmp / d.suggested_filename; await d.save_as(cv_pdf)
            check('PDF file name has the name and the employer', d.suggested_filename == 'cv-eva-lindqvist-acme-robotics.pdf', d.suggested_filename)
            raw_pdf = cv_pdf.read_bytes()
            check('the download is a real PDF', raw_pdf.startswith(b'%PDF-1.4') and raw_pdf.rstrip().endswith(b'%%EOF'))
            import pypdf
            reader = pypdf.PdfReader(str(cv_pdf)); pdf_text = ' '.join(' '.join(x.extract_text() for x in reader.pages).split())
            check('the PDF holds the whole CV, with Swedish letters', all(x in pdf_text for x in ['Eva Lindqvist', 'NYCKELKOMPETENSER', 'ERFARENHET', 'Kvalitetsingenjör, Nordvik Automotive AB, Göteborg', '2018 - nu',
                  'Produktionstekniker, Svea Bearings AB, Göteborg', 'Högskoleingenjör i maskinteknik, Chalmers', '2013-2016', 'Svenska (modersmål), engelska (flytande)', 'Ledde FMEA för bromssystem']), pdf_text)
            check('no Markdown, tips or notes in the PDF', not any(x in pdf_text for x in ['##', '**', 'kan tas bort', 'Skrivet av AI', 'Written by AI']))
            check('no long dashes anywhere in the CV, ranges included', not any(c in pdf_text for c in '\u2013\u2014') and '\u2013' not in await dlg.locator('.tl-out').inner_text())
            check('a lopsided CV with no buzzwords gets no style notes', await dlg.locator('.fc-style').count() == 0)
            check('the PDF has a title, an author and a language', reader.metadata.title == 'CV – Eva Lindqvist' and reader.metadata.author == 'Eva Lindqvist' and reader.trailer['/Root']['/Lang'] == 'sv-SE', reader.metadata)
            links = [a.get_object()['/A']['/URI'] for x in reader.pages for a in (x.get('/Annots') or [])]
            check('the email address in the PDF can be clicked', links == ['mailto:eva.lindqvist@example.com'], links)
            check('the CV fits on one page', len(reader.pages) == 1)
            if shutil.which('qpdf'):
                q = subprocess.run(['qpdf', '--check', str(cv_pdf)], capture_output=True, text=True)
                check('qpdf finds nothing wrong with the PDF', q.returncode == 0 and 'No syntax or stream encoding errors' in q.stdout, q.stdout + q.stderr)
            else: print('SKIPPED qpdf check: qpdf is not installed')
            if shutil.which('pdftoppm'): subprocess.run(['pdftoppm', '-r', '90', '-png', '-singlefile', str(cv_pdf), str(SHOTS / 'cv-pdf-sv')], check=True)
            async with pg.expect_download() as dl: await dlg.locator('[data-tl=word]').click()
            d = await dl.value; cv_word = tmp / d.suggested_filename; await d.save_as(cv_word)
            import docx
            cv_paras = [x.text for x in docx.Document(str(cv_word)).paragraphs]
            check('the Word CV starts with the name, with no title of ours', d.suggested_filename == 'cv-eva-lindqvist-acme-robotics.docx' and cv_paras[0] == 'Eva Lindqvist'
                  and not any('Anpassat CV' in x or 'Tailored CV' in x for x in cv_paras) and 'Kvalitetsingenjör, Nordvik Automotive AB, Göteborg (2018 - nu)' in cv_paras, cv_paras[:4])
            await dlg.locator('.tl-result .row').last.scroll_into_view_if_needed()
            await pg.screenshot(path=str(SHOTS / 'cv-en.png'))

            # Edit text: what the person changes is checked again in the browser, with the server's own code
            await dlg.locator('[data-tl=edit]').click()
            ed = dlg.locator('#tl-edit')
            check('Edit text opens the text as Markdown, with focus in it', (await ed.input_value()).startswith('# Eva Lindqvist\n') and await pg.evaluate("document.activeElement.id") == 'tl-edit')
            check('no downloads while editing', await dlg.locator('[data-tl=pdf]').count() == 0 and await dlg.locator('[data-tl=save]').count() == 1)
            original = await ed.input_value()
            await ed.fill(original.replace('## Utbildning', '### Kvalitetschef, Bosch Rexroth (2010 – 2012)\n- Ledde revisioner\n\n## Utbildning'))
            await dlg.locator('[data-tl=save]').click()
            warn = await dlg.locator('.fc-warn').inner_text()
            check('an added job is flagged when the edit is saved', 'Bosch Rexroth' in warn and 'Kvalitetschef' in warn and '2010' in warn and 'Checked again after your edits' in warn, warn)
            check('the edited text is shown and saved with an edit date', 'Bosch Rexroth' in await dlg.locator('.tl-out').inner_text() and 'edited' in await dlg.locator('.tl-result .row .tl-hint').inner_text()
                  and await pg.evaluate("JSON.parse(localStorage.getItem('jf-tailor'))['101'].out.cv.edited") is not None)
            await dlg.locator('[data-tl=edit]').click()
            await dlg.locator('#tl-edit').fill(original.replace('070-123 45 67', '[telefonnummer]').replace('Van att leda riskanalys', 'Van att driva riskanalys'))
            await dlg.locator('[data-task=match]').click(); await dlg.locator('[data-task=cv]').click()
            check('switching task while editing keeps the changes', 'Van att driva riskanalys' in await dlg.locator('.tl-out').inner_text() and await dlg.locator('#tl-edit').count() == 0)
            check('after the fix the check passes again', await dlg.locator('.fc-ok').count() == 1 and 'Checked again after your edits' in await dlg.locator('.fc-ok').inner_text())
            check('a reminder to fill in [brackets] before downloading', 'Fill in the parts in [brackets]' in await dlg.locator('.tl-fill').inner_text())
            async with pg.expect_download() as dl: await dlg.locator('[data-tl=pdf]').click()
            edited_text = ' '.join(pypdf.PdfReader(await (await dl.value).path()).pages[0].extract_text().split())   # lines wrap where the PDF breaks them
            check('the PDF has the edited text', 'Van att driva riskanalys' in edited_text and 'Bosch' not in edited_text and '[telefonnummer]' in edited_text, edited_text[:300])
            await dlg.locator('[data-tl=edit]').click(); await dlg.locator('#tl-edit').fill('# Eva Lindqvist\nNågot helt annat')
            await dlg.locator('[data-tl=cancel]').click()
            check('Cancel leaves the text as it was', 'Van att driva riskanalys' in await dlg.locator('.tl-out').inner_text() and 'Något helt annat' not in await dlg.locator('.tl-out').inner_text())
            await pg.evaluate("localCheck = null")   # a browser that can't run the check
            await dlg.locator('[data-tl=edit]').click(); await dlg.locator('#tl-edit').fill(original)
            await dlg.locator('[data-tl=save]').click()
            check("edits a browser can't check say so", "couldn't be checked in this browser" in await dlg.locator('.fc-warn').inner_text())
            await pg.evaluate("localCheck = undefined")

            # Characters a PDF's built-in fonts can't draw become "?", and the panel says so
            await dlg.locator('[data-tl=edit]').click(); await dlg.locator('#tl-edit').fill(original.replace('# Eva Lindqvist', '# Eva Lindqvist 李'))
            await dlg.locator('[data-tl=save]').click()
            async with pg.expect_download() as dl: await dlg.locator('[data-tl=pdf]').click()
            await dl.value
            check("characters a PDF can't show are reported", "1 characters in this text can't be shown in a PDF" in await dlg.inner_text())

            # Job Tailor's writing rules are in the prompt, and what the model still gets wrong is caught in code
            check("the prompt carries Job Tailor's writing rules", all(x in last_ai()['body']['messages'][0]['content'] for x in ['Använd inte tankstreck', 'inga modeord', 'blir aldrig "ledde"'])
                  and all(x in ask for x in ['Vikta tjänsterna ojämnt', 'Ge aldrig alla tjänster lika många punkter', 'Börja inte varje punkt med ett verb i dåtid']))
            post(FAKE + '/mode', {'tone': True})
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .fc-style', timeout=10000)
            shown = await dlg.locator('.tl-out').inner_text()
            check('dashes the model used are taken out before the CV is shown', '\u2014' not in shown and '\u2013' not in shown and 'Kvalitetsingenjör, drivande' in shown, shown[:300])
            note = await dlg.locator('.fc-style').inner_text()
            check('jobs of even length and buzzwords are pointed out', 'Every job is about the same length' in note and 'drivande' in note and 'resultatinriktad' in note, note)
            check('the style note is a status, not an alarm', await dlg.locator('.fc-style').get_attribute('role') == 'status')
            post(FAKE + '/mode', {'tone': False})
            await dlg.locator('[data-tl=edit]').click()
            fixed = (await dlg.locator('#tl-edit').input_value()).replace('Kvalitetsingenjör, drivande och resultatinriktad', 'Kvalitetsingenjör \u2014 van vid FMEA')
            fixed = fixed.replace('### Produktionstekniker', '### Produktionstekniker').split('### Produktionstekniker')[0] + '### Produktionstekniker, Svea Bearings AB, Göteborg (2016-2018)\n'
            await dlg.locator('#tl-edit').fill(fixed); await dlg.locator('[data-tl=save]').click()
            note = await dlg.locator('.fc-style').inner_text()
            check('after an edit the notes follow the text: buzzwords gone, a typed dash noted', 'drivande' not in note and 'same length' not in note and 'A dash is used as punctuation' in note, note)
            check('the dash the person typed is kept, not changed behind their back', '\u2014' in await dlg.locator('.tl-out').inner_text())

            # An answer the model stopped mid-way is marked as cut off
            post(FAKE + '/mode', {'cut': True})
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .tl-out >> nth=0', timeout=10000)
            await pg.wait_for_function("!document.querySelector('#tailor [data-tl=go]').disabled", timeout=10000)
            check('a cut-off answer says the end may be missing', 'stopped before the end' in await dlg.locator('.tl-result').inner_text() and 'Chalmers' not in await dlg.locator('.tl-out').inner_text())
            check('the new answer replaced the edited one', await pg.evaluate("JSON.parse(localStorage.getItem('jf-tailor'))['101'].out.cv.edited") is None)
            post(FAKE + '/mode', {'cut': False})

            # A tailored CV that names a job the CV doesn't have is flagged, by code on the server
            post(FAKE + '/mode', {'invent': True})
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .fc-warn', timeout=10000)
            warn = await dlg.locator('.fc-warn').inner_text()
            check('invented employer, title and date are listed above the answer', 'Fact check: 3 things here are not in your CV' in warn
                  and 'Bosch Rexroth' in warn and 'Senior kvalitetsingenjör' in warn and '2012' in warn and 'not by the AI' in warn, warn)
            check('the warning is announced', await dlg.locator('.fc-warn').get_attribute('role') == 'alert')
            await dlg.locator('.fc-warn').scroll_into_view_if_needed(); await pg.evaluate("document.querySelector('#tailor .fc-warn').scrollIntoView({block: 'center'})")
            await pg.screenshot(path=str(SHOTS / 'factcheck-en.png'))
            post(FAKE + '/mode', {'invent': False})
            stored = await pg.evaluate("JSON.parse(localStorage.getItem('jf-tailor'))['101'].out.cv.check")
            check('the check is saved with the answer', stored and stored['ok'] is False and len(stored['issues']) == 3, stored)
            await dlg.locator('[data-task=letter]').click()
            check('the letter shows its own check', await dlg.locator('.fc-ok').count() == 1 and await dlg.locator('.fc-warn').count() == 0)

            # Translate the ad inline, reusing the saved code without sending the CV
            await dlg.locator('.tl-close').click()
            trans_role = acme.locator('article.role[data-ad="101"]')
            await trans_role.locator('[data-translate]').click()
            await pg.wait_for_selector('article.role[data-ad="101"] .translation-text', timeout=10000)
            translation = await trans_role.locator('.translation-text').inner_text()
            check('English translation shown inline', 'Quality Engineer' in translation and 'quality improvements' in translation, translation)
            request_text = last_ai()['body']['messages'][1]['content']
            check('only ad text sent for translation', 'Erfarenhet av FMEA' in request_text and 'leverantörsrevisioner' not in request_text)
            check('translation code reused', await pg.evaluate("JSON.parse(localStorage.getItem('jf-code')).toUpperCase()") == CODE)
            await pg.locator('.lang button[data-lang=sv]').click()
            await pg.wait_for_function("document.querySelector('.lang button[data-lang=sv]').getAttribute('aria-pressed') === 'true'")
            bolt = pg.locator('article.co', has_text='Bolt Systems'); await bolt.locator('button.exp').click()
            check('no translation button on a Swedish ad in the Swedish interface', await bolt.locator('article.role[data-ad="103"] [data-translate]').count() == 0)
            sw_role = acme.locator('article.role[data-ad="102"]')
            check('Swedish interface offers Swedish translation of an English ad', await sw_role.locator('[data-translate]').inner_text() == 'Översätt annonsen till svenska')
            await sw_role.locator('[data-translate]').click()
            await pg.wait_for_selector('article.role[data-ad="102"] .translation-text', timeout=10000)
            check('Swedish interface requests Swedish translation and sends only ad text', len(translation_payloads) == 3 and translation_payloads[-1].get('lang') == 'sv' and translation_payloads[-1].get('job') == {'text': AD_TEXT} and 'Translate the job ad into natural Swedish' in last_ai()['body']['messages'][1]['content'], repr(translation_payloads[-1:]) + ' ' + last_ai()['body']['messages'][1]['content'])
            await pg.locator('.lang button[data-lang=en]').click()
            await pg.wait_for_function("document.querySelector('.lang button[data-lang=en]').getAttribute('aria-pressed') === 'true'")
            await trans_role.locator('[data-tailor]').click(); await pg.wait_for_selector('#tailor .tl-out', timeout=10000)

            async with pg.expect_download() as dl: await dlg.locator('[data-tl=word]').click()
            d = await dl.value; word = tmp / d.suggested_filename; await d.save_as(word)
            check('Word file name has the name from the CV', d.suggested_filename == 'cover-letter-eva-lindqvist-acme-robotics.docx', d.suggested_filename)
            doc = docx.Document(str(word)); paras = [x.text for x in doc.paragraphs]
            check('Word file opens and has the letter, with no title of ours', paras[0] == 'Cover letter' and 'Dear [contact person],' in paras and '• Eight years of FMEA' in paras, paras)
            async with pg.expect_download() as dl: await dlg.locator('[data-tl=pdf]').click()
            d = await dl.value
            letter_pdf = pypdf.PdfReader(await d.path())
            check('the cover letter downloads as a PDF too', d.suggested_filename == 'cover-letter-eva-lindqvist-acme-robotics.pdf' and 'Dear [contact person],' in letter_pdf.pages[0].extract_text()
                  and letter_pdf.trailer['/Root']['/Lang'] == 'en-GB', d.suggested_filename)
            check('bold kept in Word', any(r.bold and r.text == 'Acme Robotics' for x in doc.paragraphs for r in x.runs))
            if shutil.which('soffice'):
                conv = subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(word)], capture_output=True, timeout=120)
                check('LibreOffice can open the Word file', (tmp / (word.stem + '.pdf')).exists(), conv.stderr[-300:].decode(errors='replace'))
            else: print('SKIPPED LibreOffice open check: soffice is not installed')
            await pg.screenshot(path=str(SHOTS / 'tailor-en.png'))
            await pg.evaluate("document.documentElement.dataset.palette = 'terracotta'")
            check('the panel follows the chosen colours', await dlg.locator('.tl-head').evaluate('e => getComputedStyle(e).backgroundColor') == 'rgb(156, 74, 51)')
            await pg.evaluate("delete document.documentElement.dataset.palette")

            # Applying from the panel
            check('apply button in the panel header', await dlg.locator('.tl-head [data-tl=apply]').get_attribute('href') == 'https://apply.example/101')
            async with ctx.expect_page() as newpage: await dlg.locator('.tl-result [data-tl=apply]').click()
            await (await newpage.value).close()
            await pg.wait_for_selector('#tailor [data-tl=mark]')
            check('asks whether the application was sent', 'Sent your application?' in await dlg.locator('.tl-ask').inner_text())
            await dlg.locator('[data-tl=mark]').click()
            check('mark as applied from the panel', 'Marked as applied' in await dlg.inner_text() and await acme.locator('article.role[data-ad="101"] select[data-role=status]').input_value() == 'Applied')
            async with ctx.expect_page() as newpage: await dlg.locator('.tl-result [data-tl=apply]').click()
            await (await newpage.value).close(); await pg.wait_for_timeout(100)
            check('no question once applied', await dlg.locator('[data-tl=mark]').count() == 0)

            # Errors from the AI service and the network
            post(FAKE + '/mode', {'status': 429})
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .tl-err', timeout=10000)
            check('AI busy: clear message', "couldn't answer right now" in await dlg.locator('.tl-err').inner_text())
            check('earlier answer kept after an error', await dlg.locator('.tl-out').count() == 1)
            post(FAKE + '/mode', {'status': 200})
            await pg.route(SERVER + '/**', lambda r: r.abort())
            await dlg.locator('[data-tl=go]').click(); await pg.wait_for_selector('#tailor .tl-err', timeout=10000)
            check('server down: clear message', "Couldn't reach the tailoring server" in await dlg.locator('.tl-err').inner_text())
            await pg.unroute(SERVER + '/**')

            # Closing returns focus, and everything is still there next time
            await pg.keyboard.press('Escape')
            check('Escape closes the panel', not await dlg.evaluate('d => d.open'))
            check('focus back on the button', await pg.evaluate("document.activeElement.dataset.tailor") == '101')
            check('button shows that texts exist', await btn.locator('.tl-dot').count() == 1)
            await btn.click(); await pg.wait_for_selector('#tailor .tl-out')
            check('reopens on the last task with its answer', await dlg.locator('[data-task=letter]').get_attribute('aria-selected') == 'true' and 'Dear [contact person]' in await dlg.locator('.tl-out').inner_text())
            await dlg.locator('.tl-close').click()

            # When the full ad can't be fetched, the panel says so and fetches it again on request
            AD_DOWN['on'] = True
            await bolt.locator('article.role[data-ad="103"] [data-tailor]').click()
            await pg.wait_for_selector('#tailor [data-tl=refetch]', timeout=15000)
            check('a failed fetch says only an extract is shown', "Couldn't fetch the full ad from Platsbanken just now" in await dlg.locator('.tl-box').first.inner_text())
            AD_DOWN['on'] = False
            await dlg.locator('[data-tl=refetch]').click()
            await pg.wait_for_function("document.getElementById('tl-ad').value.includes('Kvalifikationer')", timeout=10000)
            check('fetch again fills in the full ad', 'Fetched from Platsbanken' in await dlg.locator('.tl-box').first.inner_text() and await dlg.locator('[data-tl=refetch]').count() == 0)
            await pg.set_viewport_size({'width': 320, 'height': 256})
            check('at 400% zoom the panel scrolls as a whole', await dlg.evaluate("d => getComputedStyle(d).overflowY === 'auto' && d.scrollWidth <= 320"))
            await pg.set_viewport_size({'width': 1280, 'height': 1000})
            await pg.keyboard.press('Escape')

            # Ads from other sites only have a short extract
            echo = pg.locator('article.co', has_text='Echo Apps'); await echo.locator('button.exp').click()
            await echo.locator('[data-tailor]').click(); await pg.wait_for_timeout(200)
            hint, adv = await dlg.locator('.tl-hint').first.inner_text(), await pg.input_value('#tl-ad')
            check('other-site ad: apply on that site', await dlg.locator('.tl-head [data-tl=apply]').inner_text() == 'Apply on The Hub')
            check('other-site ad: asks for the full text', 'This ad is on The Hub' in hint and adv == 'Join us as a software engineer.', (hint, adv))
            await pg.keyboard.press('Escape')

            # From the applications tab, in Swedish
            await pg.click('#tab-app')
            await pg.click('.lang button[data-lang=sv]')
            await pg.locator('#apps [data-tailor="101"]').click(); await pg.wait_for_selector('#tailor .tl-out')
            check('panel in Swedish', await dlg.locator('.tl-close').get_attribute('aria-label') == 'Stäng' and await dlg.locator('[data-task=letter]').inner_text() == 'Personligt brev ●')
            await pg.set_viewport_size({'width': 375, 'height': 800}); await pg.wait_for_timeout(200)
            check('panel fits a phone', await dlg.evaluate('d => d.scrollWidth <= d.clientWidth && d.getBoundingClientRect().width <= 375'))
            check('ad and CV stacked on a phone', await pg.evaluate("document.getElementById('tl-cv').getBoundingClientRect().top > document.getElementById('tl-ad').getBoundingClientRect().bottom"))
            await pg.screenshot(path=str(SHOTS / 'tailor-phone-sv.png'))
            await pg.keyboard.press('Escape'); await pg.set_viewport_size({'width': 1280, 'height': 1000})

            # The backup carries the CV and the texts
            async with pg.expect_download() as dl: await pg.click('#bk-down')
            backup = await (await dl.value).path()
            data = json.loads(pathlib.Path(backup).read_text(encoding='utf-8'))['data']
            check('backup has the CV and texts', 'leverantörsrevisioner' in data['cv'] and 'letter' in data['tailor']['101']['out'])
            await pg.evaluate('localStorage.clear()'); await pg.reload(); await pg.click('#tab-app')
            await pg.set_input_files('#bk-file', backup); await pg.wait_for_timeout(300)
            await pg.locator('#apps [data-tailor="101"]').click(); await pg.wait_for_selector('#tailor .tl-out')
            check('restore brings back the CV and texts', 'leverantörsrevisioner' in await pg.input_value('#tl-cv') and 'Dear [contact person]' in await dlg.locator('.tl-out').inner_text())
            check('the access code is not in the backup', 'TESTA' not in json.dumps(data).upper())

            check('no script errors', not errors, errors)
            await browser.close()
    finally:
        for pr in procs: pr.terminate()
    print(f'{len(passed)} checks passed')

if __name__ == '__main__':
    try: asyncio.run(main())
    except AssertionError as e: print('FAILED', e); sys.exit(1)
